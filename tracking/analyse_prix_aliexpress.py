import json, re, statistics, math, sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.formatting.rule import CellIsRule
from openpyxl.utils import get_column_letter

S = sys.argv[1]
OUT = sys.argv[2]
cat = json.load(open(f"{S}/catalog.json"))

SIZE_KEYS = ("Taille", "Tour de poitrine")
MIN_COEF = 2.5
HARNESS_POINTS = [34.99, 39.99, 44.99, 49.99, 59.99]
ACC_POINTS = [9.99, 12.99, 14.99, 16.99, 19.99, 24.99, 29.99]

def opt(v, names):
    for o in v["selectedOptions"]:
        if o["name"] in names:
            return o["value"]
    return None

def ref_cost(p):
    vs = [v for v in p["variants"]["nodes"] if v["inventoryItem"] and v["inventoryItem"]["unitCost"]]
    # option "Expédition" : garder l'envoi standard
    std = [v for v in vs if (opt(v, ("Expédition",)) or "Standard").lower().startswith("standard")]
    vs = std or vs
    sizes = [opt(v, SIZE_KEYS) for v in vs]
    m = [v for v, s in zip(vs, sizes) if s and re.match(r"^M\b", s)]
    if m:
        pick, label = m, "taille M"
    elif any(sizes):
        uniq = sorted({s for s in sizes if s}, key=lambda s: min(float(v["inventoryItem"]["unitCost"]["amount"]) for v, x in zip(vs, sizes) if x == s))
        mid = uniq[(len(uniq) - 1) // 2]
        pick, label = [v for v, s in zip(vs, sizes) if s == mid], f"taille médiane ({mid})"
    else:
        pick, label = vs, "toutes variantes"
    costs = [float(v["inventoryItem"]["unitCost"]["amount"]) for v in pick]
    return round(statistics.median(costs), 2), label, min(float(v["inventoryItem"]["unitCost"]["amount"]) for v in vs), max(float(v["inventoryItem"]["unitCost"]["amount"]) for v in vs)

def category(t):
    tl = t.lower()
    if tl.startswith(("laisse", "distributeur", "ceinture", "attache")):
        return "accessoire"
    if "tactique" in tl:
        return "tactique"
    if any(k in tl for k in ("canicross", "sacoche", "cuir", "ergonomique", "sécurité")):
        return "premium"
    if any(k in tl for k in ("chiot", "petit chien", "macaron")) or re.fullmatch(r"harnais chien (noir|rouge|gris|rose)", tl):
        return "entrée"
    return "cœur"

TARGET = {"entrée": 34.99, "cœur": 44.99, "tactique": 39.99, "premium": 49.99}
CAP = {"entrée": 59.99, "cœur": 59.99, "tactique": 59.99, "premium": 59.99, "accessoire": 29.99}

KW = [
    ("canicross", "canicross dog harness running"), ("sacoche", "dog backpack harness saddle bag hiking"),
    ("handicapé", "dog rear lift support harness"), ("cuir", "leather dog harness"),
    ("3 points", "escape proof dog harness 3 point"), ("airtag", "airtag dog harness no pull"),
    ("tactique", "tactical dog harness molle"), ("chiot", "puppy harness mesh breathable"),
    ("petit chien", "small dog harness vest"), ("macaron", "small dog harness vest cute"),
    ("ceinture", "dog car seat belt"), ("appuie-tête", "dog car headrest seat belt"),
    ("distributeur", "dog poop bag dispenser"), ("mains libres", "hands free bungee dog leash reflective"),
    ("laisse tactique", "tactical bungee dog leash double handle"), ("laisse", "reflective dog leash padded handle"),
    ("réfléchissant", "reflective no pull dog harness padded"), ("anti traction", "no pull dog harness padded"),
]
def search_url(t):
    tl = t.lower()
    kw = next((k for key, k in KW if key in tl), "dog harness")
    return "https://fr.aliexpress.com/w/wholesale-" + kw.replace(" ", "-") + ".html", kw

FAMILIES = [
    ("Canicross réfléchissant", r"^Harnais Canicross Réfléchissant"),
    ("Randonnée sacoche", r"^Harnais Randonnée Sacoche"),
    ("Harnais Y basique", r"^Harnais Chien (Noir|Rouge|Gris|Rose)$"),
    ("Harnais Y réfléchissant", r"^Harnais Chien (Noir|Rouge|Bleu) Réfléchissant$"),
    ("Handicapé", r"Handicapé (Camouflage|Violet|Gris|Rose|Bleu|Noir)"),
]

def cost_of(ali):
    return ali + (1.99 if ali < 10 else 0) + 3.60

def recommend(price, cost, cat):
    floor = MIN_COEF * cost
    points = ACC_POINTS if cat == "accessoire" else HARNESS_POINTS
    if cat == "accessoire":
        cand = [p for p in points if p >= floor]
        if not cand:
            return None
        rec = cand[0]
        return max(rec, min(price, rec)) if price >= rec else rec
    target = TARGET[cat]
    cand = [p for p in points if p >= max(target, floor)]
    if not cand:
        return None
    return cand[0]

rows = []
for p in cat:
    if p["productType"] == "Guide numérique":
        continue
    vs = p["variants"]["nodes"]
    price = min(float(v["price"]) for v in vs)
    cmp_vals = [float(v["compareAtPrice"]) for v in vs if v["compareAtPrice"] and float(v["compareAtPrice"]) > 0]
    cmp_ = max(cmp_vals) if cmp_vals else None
    ali, label, cmin, cmax = ref_cost(p)
    c = category(p["title"])
    cost = cost_of(ali)
    rec = recommend(price, cost, c)
    url, kw = search_url(p["title"])
    coef_now = price / cost
    notes = []
    if rec is None:
        rec = CAP[c]
        coef_cap = rec / cost
        need = math.ceil(MIN_COEF * cost) - 0.01
        if coef_cap < 2.0:
            action = "RETIRER / changer de fournisseur"
            notes.append(f"Même au plafond {CAP[c]:.2f} € le coef. n'est que {coef_cap:.2f} (2,5 exigerait ≥ {need:.2f} €).")
        else:
            action = "GARDER (marge faible)" if abs(rec - price) < 0.01 else ("BAISSER (marge faible)" if rec < price else "AUGMENTER (marge faible)")
            notes.append(f"Coef. 2,5 impossible sous le plafond {CAP[c]:.2f} € (il faudrait ≥ {need:.2f} €) : renégocier le fournisseur ou en trouver un moins cher.")
    else:
        action = None
    if rec > 1.5 * price:
        action = "RETIRER / changer de fournisseur"
        notes.append(f"Prix nécessaire ({rec:.2f} €) irréaliste vs prix actuel : à retirer ou à proposer uniquement en bundle/upsell (la taxe d'import de 3,60 € pèse trop sur un petit article). Prix laissé inchangé en attendant.")
        rec = price
    if action is None:
        action = "BAISSER" if rec < price - 0.001 else ("AUGMENTER" if rec > price + 0.001 else "GARDER")
    new_coef = rec / cost
    fam = next((f for f, rx in FAMILIES if re.search(rx, p["title"])), None)
    if c == "tactique":
        notes.append("Tactique → repositionné sur 39,99–44,99 € (Rabbitgoo Amazon 27–41 €)." if rec <= 44.99 else "Tactique : marge insuffisante pour descendre à 44,99 €.")
    if c == "premium" and rec >= 59.99:
        notes.append("Premium : plafond 59,99 € (Ruffwear Front Range ~60 €).")
    if cmp_:
        disc = 1 - price / cmp_
        if disc > 0.25:
            notes.append(f"Prix barré actuel = -{disc:.0%} : hors fourchette réaliste (-15/-25 %), risque Omnibus.")
    if cmax - cmin > 0.35 * ali and cmax - cmin > 3:
        notes.append(f"Coût très variable selon la variante ({cmin:.2f}–{cmax:.2f} €) : vérifier la marge sur les grandes tailles.")
    bar = math.floor(rec / 0.80) - 0.01  # ~ -20 %
    if action in ("BAISSER",) and price >= bar:
        notes.append(f"Prix barré possible à {bar:.2f} € seulement si ce prix a réellement été pratiqué les 30 derniers jours (Omnibus).")
    rows.append(dict(title=p["title"], url=p["onlineStoreUrl"], img=(p.get("featuredMedia") or {}).get("preview", {}).get("image", {}).get("url"),
                     price=price, cmp=cmp_, ali=ali, label=label, cat=c, rec=rec, action=action, fam=fam, bar=bar,
                     search=url, kw=kw, notes=" ".join(notes), cost=cost, coef_now=coef_now, new_coef=new_coef))

# harmonisation : même modèle décliné par coloris -> même prix
for fam in {d["fam"] for d in rows if d["fam"]}:
    grp = [d for d in rows if d["fam"] == fam]
    top = max(d["rec"] for d in grp)
    for d in grp:
        if d["rec"] < top:
            d["rec"] = top
            d["notes"] = (d["notes"] + f" Prix aligné sur les autres coloris de la famille « {fam} ».").strip()
            d["action"] = "BAISSER" if top < d["price"] - 0.001 else ("AUGMENTER" if top > d["price"] + 0.001 else "GARDER")
        d["new_coef"] = d["rec"] / d["cost"]
        d["bar"] = math.floor(d["rec"] / 0.80) - 0.01

# ---------- Excel ----------
wb = Workbook()
ws = wb.active
ws.title = "Analyse prix"
H = ["Produit", "Prix de vente actuel", "Prix barré actuel", "Prix AliExpress (meilleure correspondance)", "Lien AliExpress",
     "Frais de port ajoutés", "Taxe", "Coût de revient", "Coefficient actuel", "Prix recommandé", "Nouveau coefficient", "Commentaire",
     "Action", "Gamme", "Prix barré conseillé (~-20 %)", "Base du prix d'achat", "Statut vérification AliExpress", "Mots-clés de recherche (EN)", "URL boutique", "URL photo principale"]
ws.append(H)
hdr_fill = PatternFill("solid", fgColor="1F3A5F")
for i, h in enumerate(H, 1):
    c = ws.cell(row=1, column=i)
    c.font = Font(bold=True, color="FFFFFF")
    c.fill = hdr_fill
    c.alignment = Alignment(wrap_text=True, vertical="center")
input_font = Font(color="0000FF")
for r, d in enumerate(rows, 2):
    ws.cell(r, 1, d["title"]).hyperlink = d["url"]
    ws.cell(r, 2, d["price"])
    ws.cell(r, 3, d["cmp"])
    ws.cell(r, 4, d["ali"]).font = input_font
    lk = ws.cell(r, 5, "Recherche à vérifier")
    lk.hyperlink = d["search"]; lk.font = Font(color="0563C1", underline="single")
    ws.cell(r, 6, f"=IF(D{r}<10,1.99,0)")
    ws.cell(r, 7, 3.60).font = input_font
    ws.cell(r, 8, f"=D{r}+F{r}+G{r}")
    ws.cell(r, 9, f"=B{r}/H{r}")
    ws.cell(r, 10, d["rec"])
    ws.cell(r, 11, f"=J{r}/H{r}")
    ws.cell(r, 12, d["notes"])
    ws.cell(r, 13, d["action"])
    ws.cell(r, 14, d["cat"])
    ws.cell(r, 15, d["bar"])
    ws.cell(r, 16, f"Coût par article Shopify — {d['label']}")
    ws.cell(r, 17, "NON VÉRIFIÉ (accès AliExpress bloqué)")
    ws.cell(r, 18, d["kw"])
    ws.cell(r, 19, d["url"])
    ws.cell(r, 20, d["img"])
    for col in (2, 3, 4, 6, 7, 8, 10, 15):
        ws.cell(r, col).number_format = '#,##0.00 "€"'
    for col in (9, 11):
        ws.cell(r, col).number_format = "0.00"
    ws.cell(r, 12).alignment = Alignment(wrap_text=True, vertical="top")
last = len(rows) + 1
red = PatternFill("solid", fgColor="F8CBAD"); orange = PatternFill("solid", fgColor="FFE699"); green = PatternFill("solid", fgColor="C6EFCE")
for col in ("I", "K"):
    rng = f"{col}2:{col}{last}"
    ws.conditional_formatting.add(rng, CellIsRule(operator="lessThan", formula=["2"], fill=red))
    ws.conditional_formatting.add(rng, CellIsRule(operator="between", formula=["2", "2.4999"], fill=orange))
    ws.conditional_formatting.add(rng, CellIsRule(operator="greaterThanOrEqual", formula=["2.5"], fill=green))
widths = [42, 11, 11, 14, 18, 10, 8, 11, 10, 11, 10, 60, 16, 11, 12, 34, 22, 34, 40, 40]
for i, w in enumerate(widths, 1):
    ws.column_dimensions[get_column_letter(i)].width = w
ws.row_dimensions[1].height = 45
ws.freeze_panes = "B2"
ws.auto_filter.ref = f"A1:{get_column_letter(len(H))}{last}"

# ---------- Résumé ----------
sm = wb.create_sheet("Résumé")
bold = Font(bold=True)
sm["A1"] = "Résumé — analyse des prix Harnais Chien Expert"; sm["A1"].font = Font(bold=True, size=14)
sm["A2"] = "Prix d'achat = « coût par article » enregistré dans Shopify (variante taille M), PAS encore vérifié sur AliExpress. Modifier la colonne D de l'onglet « Analyse prix » : coûts et coefficients se recalculent."
sm["A2"].alignment = Alignment(wrap_text=True); sm.merge_cells("A2:F2"); sm.row_dimensions[2].height = 45
sm["A4"] = "Nombre de produits analysés"; sm["B4"] = len(rows)
sm["A5"] = "Coefficient moyen actuel"; sm["B5"] = f"=AVERAGE('Analyse prix'!I2:I{last})"
sm["A6"] = "Coefficient moyen après recommandation"; sm["B6"] = f"=AVERAGE('Analyse prix'!K2:K{last})"
sm["A7"] = "Produits sous le coefficient 2,5 (actuel)"; sm["B7"] = f"=COUNTIF('Analyse prix'!I2:I{last},\"<2.5\")"
sm["A8"] = "Produits sous le coefficient 2,5 (après)"; sm["B8"] = f"=COUNTIF('Analyse prix'!K2:K{last},\"<2.5\")"
for c in ("B5", "B6"):
    sm[c].number_format = "0.00"
for r in range(4, 9):
    sm.cell(r, 1).font = bold

def block(start, title, items, cols):
    sm.cell(start, 1, title).font = Font(bold=True, size=12)
    for i, h in enumerate(cols, 1):
        c = sm.cell(start + 1, i, h); c.font = Font(bold=True, color="FFFFFF"); c.fill = hdr_fill
    for j, it in enumerate(items):
        for i, v in enumerate(it, 1):
            c = sm.cell(start + 2 + j, i, v)
            if isinstance(v, float):
                c.number_format = "0.00" if v < 10 and i in (4, 6) else '#,##0.00 "€"'
    return start + 3 + len(items)

srt = sorted(rows, key=lambda d: d["coef_now"], reverse=True)
cols = ["Produit", "Prix actuel", "Coût de revient", "Coefficient actuel", "Prix recommandé", "Nouveau coef."]
fmt = lambda d: [d["title"], d["price"], round(d["cost"], 2), round(d["coef_now"], 2), d["rec"], round(d["new_coef"], 2)]
r = block(10, "Top 5 — produits les plus margés", [fmt(d) for d in srt[:5]], cols)
r = block(r + 1, "Top 5 — produits les moins margés", [fmt(d) for d in srt[::-1][:5]], cols)
for act, keys in (("À BAISSER", ("BAISSER",)), ("À GARDER", ("GARDER",)), ("À AUGMENTER", ("AUGMENTER",)), ("MARGE FAIBLE — renégocier le fournisseur", ("(marge faible)",)), ("À RETIRER / changer de fournisseur", ("RETIRER",))):
    items = [fmt(d) for d in rows if (d["action"] in keys) or (keys == ("(marge faible)",) and "marge faible" in d["action"]) or (keys == ("RETIRER",) and d["action"].startswith("RETIRER"))]
    r = block(r + 1, f"{act} ({len(items)})", items, cols)
sm.column_dimensions["A"].width = 52
for c in "BCDEF":
    sm.column_dimensions[c].width = 16

# ---------- Méthode ----------
me = wb.create_sheet("Méthode")
lines = [
    "MÉTHODE ET LIMITES",
    "1. Catalogue : lu en lecture seule via l'API Admin Shopify (produits actifs, PDF offert exclu). Aucun prix n'a été modifié.",
    "2. Prix d'achat : l'accès à fr.aliexpress.com et à harnais-chien-expert.fr est bloqué par la politique réseau de l'environnement cloud → recherche par image impossible.",
    "   À la place : « coût par article » de chaque variante dans Shopify (renseigné à l'import fournisseur). Variante de référence : taille M, sinon taille médiane ; médiane des coloris.",
    "   Ces coûts ne sont PAS vérifiés : ils peuvent dater, inclure ou non la livraison. Colonne « Lien AliExpress » = lien de recherche par mots-clés à contrôler, pas une annonce validée.",
    "3. Coût de revient = prix AliExpress + 1,99 € si < 10 € + 3,60 € de taxe d'import. Coefficient = prix TTC / coût de revient.",
    "4. Prix recommandé : palier le plus bas qui respecte coef ≥ 2,5 ET le positionnement de gamme.",
    "   Paliers harnais : 34,99 / 39,99 / 44,99 / 49,99 / 59,99 €. Gammes : entrée (chiot, petit chien, Y basique) ≥ 34,99 ; cœur ≥ 44,99 ; tactique ≥ 39,99 ; premium (canicross, sacoche, cuir, ergonomique, sécurité) ≥ 49,99 ; plafond 59,99.",
    "   Accessoires (laisses, ceintures, distributeur) : paliers 9,99 → 29,99 €, plafond 29,99 €.",
    "   MARGE FAIBLE = coef 2,5 impossible sous le plafond mais ≥ 2,0 → prix au plafond, renégocier/changer de fournisseur.",
    "   RETIRER = coef < 2,0 même au plafond, ou hausse nécessaire > +50 % → retirer, changer de fournisseur ou vendre en bundle.",
    "   Les modèles déclinés par coloris (canicross réfléchissant, sacoche, Y basique, Y réfléchissant, handicapé) reçoivent le même prix.",
    "5. Prix barré conseillé ≈ -20 %. Directive Omnibus : le prix barré doit être le prix le plus bas pratiqué dans les 30 jours précédant la réduction.",
    "   Ne pas afficher un prix barré inventé ; en cas de baisse de prix, le prix actuel ne peut servir de référence que s'il a réellement été pratiqué.",
    "Repères concurrents : Ruffwear Front Range ~60 € · Julius-K9 IDC dès ~24 € · Truelove 36–80 € · Rabbitgoo (Amazon) 27–41 € · Dog Copenhagen 38–50 €.",
]
for i, l in enumerate(lines, 1):
    me.cell(i, 1, l)
me["A1"].font = Font(bold=True, size=13)
me.column_dimensions["A"].width = 160

wb.save(OUT)
from collections import Counter
print(Counter(d["action"] for d in rows), len(rows))
print("coef moyen", round(statistics.mean(d["coef_now"] for d in rows), 2), "->", round(statistics.mean(d["new_coef"] for d in rows), 2))
for d in rows:
    print(f'{d["title"][:40]:40} {d["cat"]:10} {d["price"]:6} ali {d["ali"]:6} cost {d["cost"]:6.2f} coef {d["coef_now"]:.2f} -> {d["rec"]} ({d["new_coef"]:.2f}) {d["action"]}')
